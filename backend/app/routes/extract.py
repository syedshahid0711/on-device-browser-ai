import io
from fastapi import APIRouter, UploadFile, File, HTTPException
from app.services.llm_service import extract_profile_from_text
from app.utils.logger import get_logger
import docx

log = get_logger("extract")
router = APIRouter(prefix="/api/extract", tags=["Extract"])

@router.post("/profile")
async def extract_profile(file: UploadFile = File(...)):
    """
    Accepts a .docx file, extracts raw text with advanced formatting, and uses the local LLM
    to pull out profile fields into a JSON structure.
    """
    if not file.filename.endswith(".docx"):
        raise HTTPException(status_code=400, detail="Only .docx files are supported")
    
    try:
        content = await file.read()
        doc = docx.Document(io.BytesIO(content))
        
        text_blocks = []
        
        # Advanced extraction: Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_blocks.append(para.text.strip())
        
        # Advanced extraction: Extract tables with clear visual boundaries
        for table in doc.tables:
            text_blocks.append("\n[TABLE START]")
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                # Filter empty rows
                if any(row_data):
                    text_blocks.append(" | ".join(row_data))
            text_blocks.append("[TABLE END]\n")
                    
        full_text = "\n".join(text_blocks)
        
        if not full_text.strip():
            raise HTTPException(status_code=400, detail="The document is empty or contains no readable text")
            
        log.info(f"Extracted {len(full_text)} characters from {file.filename}")
        
        # Pass the raw text to Ollama for structured extraction
        extracted_data = await extract_profile_from_text(full_text)
        
        return {"success": True, "data": extracted_data}
        
    except Exception as e:
        log.error(f"Error extracting profile: {e}")
        raise HTTPException(status_code=500, detail=str(e))
